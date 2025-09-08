#!/usr/bin/env python3
"""
Extract content from Word document
"""

try:
    from docx import Document
    import sys
    import os
    
    def extract_docx_content(file_path):
        """Extract text content from a Word document"""
        try:
            doc = Document(file_path)
            content = []
            
            # Extract paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    content.append(paragraph.text)
            
            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        content.append(" | ".join(row_text))
            
            return "\n".join(content)
            
        except Exception as e:
            return f"Error reading document: {str(e)}"
    
    if __name__ == "__main__":
        file_path = "/workspace/user_input_files/WatchLockAI_v1_FULL_UPDATED.docx"
        
        if os.path.exists(file_path):
            content = extract_docx_content(file_path)
            
            # Save extracted content
            output_path = "/workspace/extract/watchlock_requirements.txt"
            os.makedirs(os.path.dirname(output_path), exist_ok=True)
            
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(content)
            
            print(f"Content extracted and saved to: {output_path}")
            print(f"Content length: {len(content)} characters")
            
            # Print first 2000 characters for preview
            print("\n=== DOCUMENT PREVIEW ===")
            print(content[:2000])
            if len(content) > 2000:
                print("\n... (content truncated)")
                
        else:
            print(f"File not found: {file_path}")

except ImportError:
    print("python-docx library not installed. Installing...")
    import subprocess
    subprocess.run(["uv", "add", "-q", "python-docx"], check=True)
    print("Library installed. Re-running extraction...")
    
    # Re-import and run
    from docx import Document
    import sys
    import os
    
    def extract_docx_content(file_path):
        """Extract text content from a Word document"""
        try:
            doc = Document(file_path)
            content = []
            
            # Extract paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    content.append(paragraph.text)
            
            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        content.append(" | ".join(row_text))
            
            return "\n".join(content)
            
        except Exception as e:
            return f"Error reading document: {str(e)}"
    
    file_path = "/workspace/user_input_files/WatchLockAI_v1_FULL_UPDATED.docx"
    
    if os.path.exists(file_path):
        content = extract_docx_content(file_path)
        
        # Save extracted content
        output_path = "/workspace/extract/watchlock_requirements.txt"
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(content)
        
        print(f"Content extracted and saved to: {output_path}")
        print(f"Content length: {len(content)} characters")
        
        # Print first 2000 characters for preview
        print("\n=== DOCUMENT PREVIEW ===")
        print(content[:2000])
        if len(content) > 2000:
            print("\n... (content truncated)")
            
    else:
        print(f"File not found: {file_path}")